import io
import logging
import traceback
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_file(filename: str, filepath: str) -> str:
    name = filename.lower()
    
    # 1. Try anydoc (primary parser)
    try:
        import anydoc
        logger.info(f"Using anydoc to parse {filename}")
        markdown = anydoc.to_markdown(filepath)
        if markdown and markdown.strip():
            return markdown
    except Exception as e:
        logger.warning(f"anydoc failed or unavailable for {filename}: {e}. Falling back to legacy parsers.")
        logger.debug(traceback.format_exc())
    
    # 2. Legacy fallback parsers
    try:
        if name.endswith(".pdf"):
            import pdfplumber

            text = []
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    t = page.extract_text() or ""
                    text.append(t)
            return "\n".join(text)
        elif name.endswith(".docx"):
            from docx import Document as DocxDocument

            doc = DocxDocument(filepath)
            paras = [p.text for p in doc.paragraphs]
            for tbl in doc.tables:
                for row in tbl.rows:
                    paras.append(" | ".join(c.text for c in row.cells))
            return "\n".join(paras)
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(filepath, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"=== Sheet: {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    row_txt = " | ".join("" if v is None else str(v) for v in row)
                    if row_txt.strip():
                        lines.append(row_txt)
            return "\n".join(lines)
        elif name.endswith(".pptx") or name.endswith(".ppt"):
            from pptx import Presentation
            
            prs = Presentation(filepath)
            text_runs = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_runs.append(shape.text)
            return "\n".join(text_runs)
        elif name.endswith(".txt") or name.endswith(".csv"):
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception as e:
        logger.error(f"Extraction failed for {filename}: {e}")
    return ""
